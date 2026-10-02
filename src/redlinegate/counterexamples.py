"""Eight jobs that must be refused - each a real job from the seed, changed in exactly one way,
and each caught by exactly the rule in its name."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document

from . import loop, redline, seed

DESCRIPTIONS = {
    "g1-empty-extraction": "extraction reported success with no text, and the review ran anyway",
    "g2-forced-anchor": "an edit applied although its anchor is a paraphrase that isn't in the file",
    "g3-unverified-rule": "an edit justified by a rule the pack marks unverified",
    "g4-invented-hours": "a placeholder 'helpfully' filled with hours nobody supplied",
    "g5-skipped-item": "the per-item checks stop before the last item",
    "g6-rebuilt-file": "the file rebuilt instead of edited, and the schedule table didn't survive",
    "g7-edited-after-approval": "the delivered file was changed after the requester approved it",
    "g8-feedback-dropped": "round three quietly drops a change the requester asked for in round one",
}


def _j(p):
    return json.loads(Path(p).read_text())


def _w(p, obj):
    Path(p).write_text(json.dumps(obj, indent=1) + "\n")


def _restamp(job: Path, n: int):
    """After a file changes, make the card, the shared link and the delivery agree with it, so
    the only thing wrong with the job is the thing the counterexample is about."""
    rd = job / f"round-{n}"
    card = _j(rd / "card.json")
    out_file = rd / Path(card["file"]).name
    card["file_sha"] = loop.sha(out_file)
    _w(rd / "card.json", card)
    shutil.copy(out_file, job / card["file"])
    em = job / "delivery" / "email.json"
    if em.exists() and _j(em)["approved_round"] == n:
        shutil.copy(out_file, job / "delivery" / out_file.name)
        e = _j(em)
        e["attachments"][0]["sha256"] = card["file_sha"]
        _w(em, e)


def _rebuild(job: Path, n: int, edits: list[dict]):
    rd = job / f"round-{n}"
    review, card = _j(rd / "review.json"), _j(rd / "card.json")
    review["edits"] = edits
    _w(rd / "review.json", review)
    meta = _j(job / "job.json")
    pl = redline.apply(job / meta["original"], edits, rd / Path(card["file"]).name, f"{review['today']}T00:00:00Z")
    _w(rd / "placements.json", [p.__dict__ for p in pl])
    _restamp(job, n)


def build(data: Path, out: Path) -> list[str]:
    data, out = Path(data), Path(out)
    shutil.rmtree(out, ignore_errors=True)
    seed.write(data)
    base = out / "_base"
    for job in ("deck-ashcombe", "solar-valcourt", "sign-lindqvist"):
        loop.run_job(data, job, base)

    def copy(src, name):
        d = out / name
        shutil.copytree(base / src, d)
        return d

    d = copy("sign-lindqvist", "g1-empty-extraction")
    (d / "extracted.md").write_text("")
    m = _j(d / "job.json"); m["extraction"]["chars"] = 0; _w(d / "job.json", m)

    d = copy("sign-lindqvist", "g2-forced-anchor")
    r = _j(d / "round-1" / "review.json")
    r["edits"][0]["anchor"] = "The wall sign is lit from inside."
    _w(d / "round-1" / "review.json", r)

    d = copy("deck-ashcombe", "g3-unverified-rule")
    r = _j(d / "round-1" / "review.json")
    r["edits"][1]["rule"] = "DECK-DOC-006"
    _w(d / "round-1" / "review.json", r)

    d = copy("sign-lindqvist", "g4-invented-hours")
    e = _j(d / "round-1" / "review.json")["edits"]
    e[0]["text"] = "The wall sign is internally lit, on a timer set to 6 am to 10 pm."
    _rebuild(d, 1, e)

    d = copy("deck-ashcombe", "g5-skipped-item")
    r = _j(d / "round-3" / "review.json")
    r["item_checks"] = "\n".join(r["item_checks"].splitlines()[:2])
    _w(d / "round-3" / "review.json", r)

    d = copy("solar-valcourt", "g6-rebuilt-file")
    f = d / "round-1" / Path(_j(d / "round-1" / "card.json")["file"]).name
    doc = Document(f)
    tbl = doc.element.body.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl")
    tbl.getparent().remove(tbl)
    doc.save(f)
    _restamp(d, 1)

    d = copy("sign-lindqvist", "g7-edited-after-approval")
    f = next((d / "delivery").glob("*.docx"))
    doc = Document(f)
    doc.add_paragraph("Illumination hours: dusk to close.")
    doc.save(f)
    e = _j(d / "delivery" / "email.json"); e["attachments"][0]["sha256"] = loop.sha(f); _w(d / "delivery" / "email.json", e)

    d = copy("deck-ashcombe", "g8-feedback-dropped")
    e = _j(d / "round-3" / "review.json")["edits"]
    _rebuild(d, 3, [x for x in e if not x["anchor"].startswith("the setback")])

    shutil.rmtree(base)
    return sorted(DESCRIPTIONS)
